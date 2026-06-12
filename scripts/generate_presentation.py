"""
generate_presentation.py
------------------------
Captures screenshots of the running Portfolio Dashboard (localhost:8501)
by scrolling through each tab and stitching into a single merged full-page
image, then assembles a PowerPoint presentation.

Rules:
  - Options tab  : single viewport screenshot (fits in one shot, no scroll)
  - Performance  : stitch full page, then split into 2 slides (too tall)
  - All other tabs: stitch full page into 1 slide

Usage:
    1. Start the dashboard:  streamlit run app.py
    2. Run this script:      python -X utf8 scripts/generate_presentation.py
    3. Open output:          output/Portfolio_Dashboard_Presentation.pptx

Dependencies:
    pip install python-pptx playwright Pillow
    python -m playwright install chromium
"""

import sys
import time
import io
from pathlib import Path
from datetime import date

# Force UTF-8 output on Windows terminals
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

# ---------------------------------------------------------------------------
# Dependency checks
# ---------------------------------------------------------------------------
try:
    from playwright.sync_api import sync_playwright
except ImportError:
    sys.exit("Missing: pip install playwright && python -m playwright install chromium")

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    sys.exit("Missing: pip install python-pptx")

try:
    from PIL import Image as PILImage
except ImportError:
    sys.exit("Missing: pip install Pillow")

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
APP_URL       = "http://localhost:8501"
OUTPUT_PATH   = Path(__file__).parent.parent / "output" / "Portfolio_Dashboard_Presentation.pptx"
SCREENSHOT_DIR = Path(__file__).parent.parent / "output" / "screenshots"
SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)

VIEWPORT_W = 1600
VIEWPORT_H = 900

# 16:9 widescreen slide
SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# Colours matching Streamlit dark theme
BG_COLOR       = RGBColor(0x0F, 0x11, 0x17)
TITLE_BG_COLOR = RGBColor(0x16, 0x1B, 0x2E)
ACCENT_COLOR   = RGBColor(0xFF, 0x4B, 0x4B)
TEXT_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_MUTED     = RGBColor(0xA0, 0xA0, 0xB8)
BULLET_COLOR   = RGBColor(0x63, 0xB3, 0xED)

# ---------------------------------------------------------------------------
# Tab definitions
# ---------------------------------------------------------------------------
TABS = [
    (
        "Statistics", "Statistics", "📊",
        [
            "Total portfolio value, P&L, and CAGR at a glance",
            "Portfolio Composition Treemap — allocation with P&L% color coding",
            "Benchmark Comparison Table — vs S&P 500 & TA-125",
            "Top Gainers / Top Losers side-by-side ranking",
            "Currency Exposure — NIS vs USD weight breakdown",
            "Realized P&L and average holding period",
        ],
    ),
    (
        "Performance", "Performance", "📈",
        [
            "7 time-series charts: area, drawdown, indexed benchmarks, monthly returns, rolling Sharpe",
            "Invested Capital Over Time — gradient area chart",
            "Drawdown Chart — peak-to-trough underwater plot",
            "Portfolio vs S&P 500 vs TA-125 — base-100 indexed comparison",
            "US vs TASE decomposition — separate market performance lines",
            "Monthly Returns Bar — color-coded gain/loss bars",
            "Rolling 60-day Sharpe Ratio with average threshold line",
        ],
    ),
    (
        "Cash Flow", "Cash Flow", "💰",
        [
            "Cumulative Inflows / Outflows / Net — triple overlay chart",
            "Monthly Cash Flow — stacked bar with net line overlay",
            "Net Flow by Category — horizontal breakdown bar chart",
            "Inflow Composition — donut chart by category",
            "Yearly Summary Table — deposits, withdrawals, dividends, fees",
            "Transaction Detail Expander — 200 most recent with colored flow pills",
        ],
    ),
    (
        "TASE", "TASE (₪)", "🏦",
        [
            "NIS position table with avg cost, live price, P&L pills",
            "Allocation Donut Chart — market value composition",
            "P&L Bar Chart — per-position gain/loss colored bars",
            "Free cash balance card (NIS)",
            "Agorot normalization: IBI prices ÷ 100 → ₪",
            "TASE numeric IBI IDs resolved to tickers via symbol mapper",
        ],
    ),
    (
        "US", "US ($)", "🌐",
        [
            "USD position table with avg cost, live price, P&L pills",
            "Allocation Donut Chart — USD portfolio composition",
            "P&L Bar Chart — US positions gain/loss",
            "Free cash balance card (USD)",
            "Supports numeric IBI IDs mapped to US tickers",
            "Identical layout to TASE tab for easy cross-market comparison",
        ],
    ),
    (
        "Merged", "Merged (₪)", "🌍",
        [
            "All positions unified in NIS using live USD/ILS FX rate",
            "FX rate & date reference shown at top of view",
            "Full Portfolio Allocation Donut — TASE + US combined",
            "Unified P&L Bar Chart — all positions in ₪",
            "TASE tickers shown as 'TICKER (IBI_ID)' format",
            "Cost basis uses historical FX; market value uses current FX",
        ],
    ),
    (
        "Options", "Options", "📋",
        [
            "LONG / SHORT / CLOSED direction badges with color coding",
            "'Open positions only' toggle — hides expired positions",
            "'Interactive table' toggle — Pandas dataframe vs styled HTML",
            "Auto-expiry detection from option symbol date format",
            "Total Capital shown as absolute value (handles short premium)",
            "Supports both NIS and USD options in a single view",
        ],
    ),
]

# Options: single viewport shot (fits in one screen, no scroll needed)
SINGLE_SHOT_TABS = {"Options"}

# Performance: stitch full page then split into 2 slides (very tall)
SPLIT_TABS = {"Performance"}

ARCH_STEPS = [
    ("IBI Excel (.xlsx)",    "Trans_Input/Transactions_IBI.xlsx"),
    ("Excel Reader",         "src/portfolio/ingestion.py"),
    ("IBI Classifier",       "src/classifiers/ibi_classifier.py — 21 Hebrew tx types"),
    ("FX Rates",             "USD/ILS rates fetched at ingestion time"),
    ("SQLite DB",            "data/portfolio.db — transactions, prices, states"),
    ("Portfolio Builder",    "src/portfolio/builder.py — positions, realized trades"),
    ("Price Fetcher",        "Twelvedata primary → yfinance fallback → price_cache"),
    ("Dashboard",            "app.py → 7 tabs → src/dashboard/views/"),
]

# ---------------------------------------------------------------------------
# Playwright scroll helpers
# ---------------------------------------------------------------------------

def _scroll_to(page, y: int):
    """Scroll both the window and the Streamlit main container."""
    page.evaluate(f"""
        window.scrollTo(0, {y});
        var main = document.querySelector('section[data-testid="stMain"]')
                || document.querySelector('.main');
        if (main) main.scrollTop = {y};
    """)


def _scroll_top(page):
    _scroll_to(page, 0)


def _total_height(page) -> int:
    """Return the full scrollable height of the page (window or stMain, whichever is taller)."""
    return page.evaluate(
        "Math.max("
        "  document.documentElement.scrollHeight,"
        "  document.body ? document.body.scrollHeight : 0,"
        "  (document.querySelector('section[data-testid=\"stMain\"]') || {scrollHeight:0}).scrollHeight,"
        "  (document.querySelector('.main') || {scrollHeight:0}).scrollHeight"
        ")"
    )


def _actual_y(page) -> int:
    """Return the actual current scroll position."""
    return page.evaluate(
        "Math.max("
        "  window.scrollY || window.pageYOffset || 0,"
        "  (document.querySelector('section[data-testid=\"stMain\"]') || {scrollTop:0}).scrollTop,"
        "  (document.querySelector('.main') || {scrollTop:0}).scrollTop"
        ")"
    )


# ---------------------------------------------------------------------------
# Screenshot capture & stitching
# ---------------------------------------------------------------------------

def stitch_images(shots: list[tuple[int, bytes]], total_h: int) -> bytes:
    """
    Stitch viewport screenshots into one tall image without overlap.

    shots: ordered list of (scroll_y, image_bytes).
           scroll_y is the ACTUAL scroll position when the screenshot was taken.
    total_h: full page height in pixels (canvas height).
    """
    imgs = [(y, PILImage.open(io.BytesIO(data))) for y, data in shots]
    width = imgs[0][1].width

    canvas = PILImage.new("RGB", (width, total_h), (15, 17, 23))

    covered = 0  # how many rows of the canvas have been filled
    for i, (scroll_y, img) in enumerate(imgs):
        # The row within this image that corresponds to 'covered' in the canvas
        img_row_start = max(0, covered - scroll_y)

        if i < len(imgs) - 1:
            # Number of new rows: from covered up to where the next shot starts
            next_scroll_y = imgs[i + 1][0]
            new_rows = max(0, next_scroll_y - covered)
        else:
            # Last image: fill the rest of the canvas
            new_rows = total_h - covered

        new_rows = min(new_rows, img.height - img_row_start)
        if new_rows <= 0:
            continue

        crop = img.crop((0, img_row_start, width, img_row_start + new_rows))
        canvas.paste(crop, (0, covered))
        covered += new_rows

    buf = io.BytesIO()
    canvas.save(buf, format="PNG")
    return buf.getvalue()


def capture_stitched(page) -> bytes:
    """
    Scroll through the page taking viewport screenshots, then stitch them
    into one seamless full-page image.
    """
    _scroll_top(page)
    time.sleep(0.8)

    total_h = _total_height(page)
    print(f"    page height: {total_h}px")

    # If the content fits in one viewport, no stitching needed
    if total_h <= VIEWPORT_H + 50:
        return page.screenshot(full_page=False)

    shots: list[tuple[int, bytes]] = []
    request_y = 0

    while True:
        _scroll_to(page, request_y)
        time.sleep(0.7)
        actual = _actual_y(page)
        shot = page.screenshot(full_page=False)
        shots.append((actual, shot))
        print(f"    captured strip at y={actual}")

        next_y = actual + VIEWPORT_H
        if next_y >= total_h:
            break
        request_y = next_y

    # Make sure the very bottom is included
    last_y = shots[-1][0]
    if last_y + VIEWPORT_H < total_h:
        bottom_y = total_h - VIEWPORT_H
        _scroll_to(page, bottom_y)
        time.sleep(0.7)
        actual = _actual_y(page)
        shot = page.screenshot(full_page=False)
        shots.append((actual, shot))
        print(f"    final bottom strip at y={actual}")

    return stitch_images(shots, total_h)


def split_image_halves(img_bytes: bytes) -> tuple[bytes, bytes]:
    """Split image vertically into top and bottom halves."""
    img = PILImage.open(io.BytesIO(img_bytes))
    w, h = img.size
    mid = h // 2

    top = img.crop((0, 0, w, mid))
    bot = img.crop((0, mid, w, h))

    top_buf = io.BytesIO()
    top.save(top_buf, format="PNG")
    bot_buf = io.BytesIO()
    bot.save(bot_buf, format="PNG")
    return top_buf.getvalue(), bot_buf.getvalue()


def capture_screenshots() -> dict[str, bytes]:
    """
    Drive the browser through every tab, capturing:
      "overview"       → viewport screenshot of the default view
      "Options"        → single viewport screenshot (no scroll)
      other tab keys   → stitched full-page image
    """
    screenshots: dict[str, bytes] = {}

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": VIEWPORT_W, "height": VIEWPORT_H})

        print(f"  Opening {APP_URL} ...")
        page.goto(APP_URL, wait_until="networkidle", timeout=30000)
        time.sleep(3)

        # Overview (sidebar + default tab visible)
        _scroll_top(page)
        time.sleep(1)
        print("  Capturing overview ...")
        screenshots["overview"] = page.screenshot(full_page=False)

        for selector_text, display_name, _icon, _bullets in TABS:
            print(f"  Tab: {display_name} ...")
            try:
                tab_btn = page.locator(
                    f'button[data-baseweb="tab"]:has-text("{selector_text}")'
                ).first
                tab_btn.click()
                page.wait_for_load_state("networkidle")
                time.sleep(2)
                _scroll_top(page)
                time.sleep(0.6)

                if selector_text in SINGLE_SHOT_TABS:
                    screenshots[selector_text] = page.screenshot(full_page=False)
                    print(f"    single viewport shot")
                else:
                    screenshots[selector_text] = capture_stitched(page)

            except Exception as exc:
                print(f"    WARNING: {exc}")
                screenshots[selector_text] = b""

        browser.close()

    # Persist all screenshots to disk for inspection
    for name, data in screenshots.items():
        if data:
            path = SCREENSHOT_DIR / f"{name}.png"
            path.write_bytes(data)
            size_kb = len(data) // 1024
            print(f"  Saved {path.name} ({size_kb} KB)")

    return screenshots


# ---------------------------------------------------------------------------
# PPTX helpers
# ---------------------------------------------------------------------------

def _add_bg(slide, color: RGBColor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_rect(slide, left, top, width, height, fill_color: RGBColor,
              line_color=None, line_width=0):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def _add_text(slide, text, left, top, width, height,
              font_size=18, bold=False, color=TEXT_WHITE,
              align=PP_ALIGN.LEFT, word_wrap=True):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = word_wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    return txBox


def _add_image_fitted(slide, img_bytes: bytes, area_left, area_top, area_w, area_h):
    """
    Embed the image centred within the area, preserving aspect ratio.
    The image is scaled to fill as much of the area as possible without cropping.
    """
    if not img_bytes:
        _add_rect(slide, area_left, area_top, area_w, area_h, RGBColor(0x2A, 0x2A, 0x3A))
        _add_text(slide, "[ screenshot unavailable ]",
                  area_left, area_top + area_h // 2 - Inches(0.3),
                  area_w, Inches(0.6),
                  font_size=14, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
        return

    img = PILImage.open(io.BytesIO(img_bytes))
    w_px, h_px = img.size
    aspect = w_px / h_px  # width / height ratio

    # Fit by width first
    fit_w = area_w
    fit_h = int(fit_w / aspect)

    # If too tall, fit by height instead
    if fit_h > area_h:
        fit_h = area_h
        fit_w = int(fit_h * aspect)

    # Centre in area
    off_x = (area_w - fit_w) // 2
    off_y = (area_h - fit_h) // 2

    stream = io.BytesIO(img_bytes)
    slide.shapes.add_picture(stream, area_left + off_x, area_top + off_y, fit_w, fit_h)


def _add_bullets(slide, bullets: list[str], left, top, width, height):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    hdr = tf.paragraphs[0]
    hdr.alignment = PP_ALIGN.LEFT
    hr = hdr.add_run()
    hr.text = "Key Features"
    hr.font.size = Pt(14)
    hr.font.bold = True
    hr.font.color.rgb = ACCENT_COLOR

    for bullet in bullets:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(4)
        r = p.add_run()
        r.text = f"• {bullet}"
        r.font.size = Pt(11)
        r.font.color.rgb = TEXT_WHITE


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------

def _title_bar(slide, title_text: str):
    """Standard title bar shared by all content slides."""
    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), ACCENT_COLOR)
    _add_rect(slide, 0, Inches(0.08), SLIDE_W, Inches(0.7), TITLE_BG_COLOR)
    _add_text(slide, title_text,
              Inches(0.4), Inches(0.15), Inches(12.5), Inches(0.6),
              font_size=24, bold=True)
    _add_rect(slide, 0, SLIDE_H - Inches(0.08), SLIDE_W, Inches(0.08), ACCENT_COLOR)


def build_title_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, BG_COLOR)
    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), ACCENT_COLOR)

    cx, cy, cw = Inches(1.5), Inches(1.8), Inches(10.33)

    _add_text(slide, "IBI Portfolio Dashboard",
              cx, cy + Inches(0.75), cw, Inches(1.1),
              font_size=44, bold=True)
    _add_text(slide,
              "Multi-currency investment tracker  ·  TASE & US markets  ·  Built with Streamlit",
              cx, cy + Inches(1.85), cw, Inches(0.6),
              font_size=18, color=TEXT_MUTED)
    _add_rect(slide, cx, cy + Inches(2.55), cw, Inches(0.04), ACCENT_COLOR)

    stats = [("7", "Dashboard Tabs"), ("2", "Markets"), ("₪ / $", "Multi-currency"), ("SQLite", "Local DB")]
    stat_w = Inches(2.4)
    for i, (val, label) in enumerate(stats):
        sx = cx + i * stat_w
        sy = cy + Inches(2.8)
        _add_text(slide, val, sx, sy, stat_w, Inches(0.6),
                  font_size=28, bold=True, color=ACCENT_COLOR, align=PP_ALIGN.CENTER)
        _add_text(slide, label, sx, sy + Inches(0.55), stat_w, Inches(0.4),
                  font_size=11, color=TEXT_MUTED, align=PP_ALIGN.CENTER)

    _add_text(slide, f"Generated {date.today().strftime('%B %d, %Y')}",
              0, SLIDE_H - Inches(0.5), SLIDE_W, Inches(0.4),
              font_size=10, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    _add_rect(slide, 0, SLIDE_H - Inches(0.08), SLIDE_W, Inches(0.08), ACCENT_COLOR)


def build_overview_slide(prs: Presentation, img_bytes: bytes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, BG_COLOR)
    _title_bar(slide, "Dashboard Overview")

    _add_image_fitted(slide, img_bytes,
                      Inches(0.3), Inches(0.9), Inches(8.5), Inches(5.9))

    _add_bullets(slide, [
        "Twelvedata API status indicator",
        "Import Transactions (.xlsx upload)",
        "Force Re-parse Excel button",
        "Refresh Prices button",
        "Last import metadata",
        "Database path & row count",
    ], Inches(9.0), Inches(0.9), Inches(4.1), Inches(5.9))


def build_full_page_slide(prs: Presentation,
                          display_name: str, icon: str,
                          img_bytes: bytes, subtitle: str = ""):
    """
    Single slide showing a full-page stitched screenshot.
    The image is fitted (aspect-ratio preserving) into the entire slide area
    below the title bar.
    """
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, BG_COLOR)
    title = f"{icon}  {display_name}" + (f"  —  {subtitle}" if subtitle else "")
    _title_bar(slide, title)

    # Screenshot fills everything below the title bar
    area_top = Inches(0.85)
    area_h   = SLIDE_H - area_top - Inches(0.12)
    _add_image_fitted(slide, img_bytes,
                      Inches(0.15), area_top, SLIDE_W - Inches(0.3), area_h)


def build_options_slide(prs: Presentation, img_bytes: bytes, bullets: list[str]):
    """Options tab: viewport screenshot on the left + feature bullets on the right."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, BG_COLOR)
    _title_bar(slide, "📋  Options")

    _add_image_fitted(slide, img_bytes,
                      Inches(0.3), Inches(0.9), Inches(8.5), Inches(5.9))
    _add_bullets(slide, bullets, Inches(9.0), Inches(0.9), Inches(4.1), Inches(5.9))


def build_architecture_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, BG_COLOR)
    _title_bar(slide, "⚙️  Data Pipeline & Architecture")

    box_w  = Inches(11.5)
    box_h  = Inches(0.52)
    box_x  = Inches(0.9)
    gap    = Inches(0.1)
    start_y = Inches(1.05)

    row_colors = [
        RGBColor(0x1E, 0x3A, 0x5F),
        RGBColor(0x1A, 0x35, 0x4F),
        RGBColor(0x16, 0x2E, 0x44),
        RGBColor(0x12, 0x28, 0x3A),
    ]

    for i, (step, detail) in enumerate(ARCH_STEPS):
        by = start_y + i * (box_h + gap)
        _add_rect(slide, box_x, by, box_w, box_h,
                  row_colors[i % len(row_colors)],
                  line_color=ACCENT_COLOR, line_width=0.5)
        if i < len(ARCH_STEPS) - 1:
            _add_rect(slide,
                      box_x + box_w // 2 - Inches(0.04), by + box_h,
                      Inches(0.08), gap, ACCENT_COLOR)
        _add_text(slide, str(i + 1),
                  box_x + Inches(0.1), by + Inches(0.04),
                  Inches(0.35), box_h - Inches(0.08),
                  font_size=13, bold=True, color=ACCENT_COLOR)
        _add_text(slide, step,
                  box_x + Inches(0.5), by + Inches(0.04),
                  Inches(3.5), box_h - Inches(0.08),
                  font_size=13, bold=True)
        _add_text(slide, detail,
                  box_x + Inches(4.1), by + Inches(0.06),
                  Inches(7.2), box_h - Inches(0.12),
                  font_size=10, color=TEXT_MUTED)


def build_summary_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_bg(slide, BG_COLOR)
    _title_bar(slide, "Feature Summary")

    rows = [
        ("Tab",          "Charts",                               "Controls",   "Highlights"),
        ("Statistics",   "Treemap, P&L bar, comparison table",  "—",          "Gainers/Losers, Benchmarks"),
        ("Performance",  "7 charts incl. drawdown, Sharpe",     "—",          "Dual book/market value"),
        ("Cash Flow",    "Cumulative, monthly, donut",           "—",          "Yearly table, tx expander"),
        ("TASE (₪)",     "Donut, P&L bar",                      "—",          "Agorot → ₪ normalization"),
        ("US ($)",       "Donut, P&L bar",                      "—",          "Numeric IBI → ticker mapping"),
        ("Merged (₪)",   "Donut, P&L bar",                      "—",          "Live FX rate unification"),
        ("Options",      "— (table only)",                      "2 toggles",  "Direction badges, expiry detection"),
    ]

    col_x = [Inches(0.3), Inches(3.0), Inches(6.5), Inches(9.0)]
    col_w = [Inches(2.6), Inches(3.4), Inches(2.4), Inches(4.0)]
    row_h  = Inches(0.52)
    start_y = Inches(0.95)

    for r_i, row in enumerate(rows):
        ry = start_y + r_i * row_h
        bg = TITLE_BG_COLOR if r_i == 0 else (
            RGBColor(0x1A, 0x20, 0x2E) if r_i % 2 == 0 else RGBColor(0x14, 0x19, 0x26)
        )
        _add_rect(slide, Inches(0.3), ry, Inches(12.73), row_h, bg)
        for c_i, cell in enumerate(row):
            bold  = r_i == 0
            color = ACCENT_COLOR if r_i == 0 else (BULLET_COLOR if c_i == 0 else TEXT_WHITE)
            _add_text(slide, cell,
                      col_x[c_i] + Inches(0.08), ry + Inches(0.08),
                      col_w[c_i] - Inches(0.1), row_h - Inches(0.1),
                      font_size=11, bold=bold, color=color)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    print("=" * 60)
    print("Portfolio Dashboard — Presentation Generator")
    print("=" * 60)

    # 1. Capture
    print("\n[1/3] Capturing screenshots ...")
    screenshots = capture_screenshots()

    # 2. Build PPTX
    print("\n[2/3] Building PowerPoint presentation ...")
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    build_title_slide(prs)
    print("  Title")

    build_overview_slide(prs, screenshots.get("overview", b""))
    print("  Overview / Sidebar")

    for selector_text, display_name, icon, bullets in TABS:
        img = screenshots.get(selector_text, b"")

        if selector_text in SINGLE_SHOT_TABS:
            # Options: viewport shot + feature bullets
            build_options_slide(prs, img, bullets)
            print(f"  {display_name}: 1 slide (viewport)")

        elif selector_text in SPLIT_TABS:
            # Performance: split stitched image into 2 slides
            if img:
                top_img, bot_img = split_image_halves(img)
            else:
                top_img = bot_img = b""
            build_full_page_slide(prs, display_name, icon, top_img, "Part 1 / 2")
            build_full_page_slide(prs, display_name, icon, bot_img, "Part 2 / 2")
            print(f"  {display_name}: 2 slides (split full page)")

        else:
            # All other tabs: stitched full page in one slide
            build_full_page_slide(prs, display_name, icon, img)
            print(f"  {display_name}: 1 slide (full page)")

    build_architecture_slide(prs)
    print("  Architecture")

    build_summary_slide(prs)
    print("  Summary")

    # 3. Save
    print(f"\n[3/3] Saving to {OUTPUT_PATH} ...")
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT_PATH))

    print(f"\nDone!  Open: {OUTPUT_PATH}")
    print(f"Screenshots saved in: {SCREENSHOT_DIR}")


if __name__ == "__main__":
    main()
